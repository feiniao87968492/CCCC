#!/usr/bin/env python3
"""Fill a mapped DOCX template from deterministic Markdown sections."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document


ANCHOR_KEYS = (
    "title_anchor_para_idx",
    "abstract_anchor_para_idx",
    "body_anchor_para_idx",
    "references_anchor_para_idx",
    "appendix_anchor_para_idx",
)
REQUIRED_ANCHORS = ANCHOR_KEYS[:4]
ABSTRACT_NAMES = {"摘要", "abstract"}
REFERENCE_NAMES = {"参考文献", "references", "bibliography"}
APPENDIX_NAMES = {"附录", "appendix", "annex"}


@dataclass
class MarkdownPaper:
    title: str = ""
    abstract: list[str] = field(default_factory=list)
    keywords: str = ""
    body: list[tuple[str, str | None]] = field(default_factory=list)
    references: list[str] = field(default_factory=list)
    appendix: list[tuple[str, str | None]] = field(default_factory=list)


def _safe_path(value: str, workspace: Path, *, must_exist: bool, label: str) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = workspace / path
    unresolved = path.absolute()
    resolved = path.resolve(strict=False)
    try:
        resolved.relative_to(workspace)
    except ValueError as exc:
        raise ValueError(f"{label} escapes the workspace") from exc
    current = unresolved
    while current != workspace and current.parent != current:
        if current.exists() and current.is_symlink():
            raise ValueError(f"{label} cannot use symlinks")
        current = current.parent
    if must_exist and not resolved.is_file():
        raise ValueError(f"{label} must be an existing file")
    return resolved


def _section_name(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().lower())


def _flush(buffer: list[str], target: list, style: str | None = None) -> None:
    text = " ".join(item.strip() for item in buffer if item.strip()).strip()
    if text:
        target.append((text, style) if style is not None else text)
    buffer.clear()


def parse_markdown(text: str) -> MarkdownPaper:
    paper = MarkdownPaper()
    section = "body"
    buffer: list[str] = []

    def active_target():
        if section == "abstract":
            return paper.abstract
        if section == "references":
            return paper.references
        if section == "appendix":
            return paper.appendix
        return paper.body

    def flush() -> None:
        target = active_target()
        if section in {"body", "appendix"}:
            _flush(buffer, target, "Normal")
        else:
            _flush(buffer, target)

    for raw in text.splitlines():
        line = raw.strip()
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            normalized = _section_name(title)
            if level == 1 and not paper.title:
                paper.title = title
                continue
            if normalized in ABSTRACT_NAMES:
                section = "abstract"
            elif normalized in REFERENCE_NAMES:
                section = "references"
            elif normalized in APPENDIX_NAMES:
                section = "appendix"
            else:
                section = "body"
                paper.body.append((title, f"Heading {min(level, 9)}"))
            continue
        if not line:
            flush()
            continue
        if section == "abstract" and re.match(r"^(关键词|关键字|keywords?)\s*[:：]", line, re.I):
            flush()
            paper.keywords = line
            continue
        list_match = re.match(r"^([-*+]|\d+[.)])\s+(.+)$", line)
        if list_match:
            flush()
            marker, item = list_match.groups()
            list_style = "List Number" if marker[0].isdigit() else "List Bullet"
            if section in {"body", "appendix"}:
                active_target().append((item.strip(), list_style))
            else:
                active_target().append(line)
            continue
        if section == "references" and re.match(r"^\[?\d+\]?", line):
            flush()
            paper.references.append(line)
            continue
        buffer.append(line)
    flush()
    if not paper.title:
        raise ValueError("source Markdown requires a level-1 title")
    return paper


def _int_index(value: object, key: str, limit: int, *, optional: bool = False) -> int | None:
    if optional and value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise ValueError(f"{key} must be an integer")
    if value < 0 or value >= limit:
        raise ValueError(f"{key} is out of range")
    return value


def validate_map(mapping: object, paragraph_count: int, table_count: int) -> dict[str, object]:
    if not isinstance(mapping, dict):
        raise ValueError("template map must be a JSON object")
    normalized = dict(mapping)
    anchors: dict[str, int | None] = {}
    for key in ANCHOR_KEYS:
        if key in REQUIRED_ANCHORS and key not in mapping:
            raise ValueError(f"template map is missing {key}")
        anchors[key] = _int_index(
            mapping.get(key), key, paragraph_count, optional=key == "appendix_anchor_para_idx"
        )
    present_anchors = [value for value in anchors.values() if value is not None]
    if len(present_anchors) != len(set(present_anchors)):
        raise ValueError("template anchors must be unique")

    delete_raw = mapping.get("delete_paragraph_indices", [])
    preserve_raw = mapping.get("preserve_table_indices", [])
    if not isinstance(delete_raw, list) or not isinstance(preserve_raw, list):
        raise ValueError("delete and preserve indices must be lists")
    delete = {_int_index(value, "delete paragraph index", paragraph_count) for value in delete_raw}
    preserve = {_int_index(value, "preserve table index", table_count) for value in preserve_raw}
    if set(present_anchors) & delete:
        raise ValueError("anchor indices conflict with delete_paragraph_indices")
    mode = str(mapping.get("body_anchor_mode", "delete")).strip().lower()
    if mode not in {"delete", "keep", "replace"}:
        raise ValueError("body_anchor_mode must be delete, keep, or replace")
    normalized.update(anchors)
    normalized["delete_paragraph_indices"] = sorted(delete)
    normalized["preserve_table_indices"] = sorted(preserve)
    normalized["body_anchor_mode"] = mode
    return normalized


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _style_name(paragraph) -> str | None:
    return paragraph.style.name if paragraph.style is not None else None


def _sample_style(
    originals: list,
    start: int,
    end: int,
    delete_indices: set[int],
    fallback: str | None,
) -> str | None:
    for index in range(start + 1, end):
        if index in delete_indices and originals[index].text.strip():
            return _style_name(originals[index]) or fallback
    return fallback


def _insert_elements(
    document,
    anchor,
    elements: list[tuple[str, str | None]],
    *,
    normal_style: str | None,
    heading_style: str | None,
) -> int:
    cursor = anchor._p
    count = 0
    style_names = {style.name for style in document.styles}
    for text, style in elements:
        if not text:
            continue
        if style and style.startswith("Heading"):
            requested = style if style in style_names else heading_style
        elif style in {"List Bullet", "List Number"}:
            requested = style
        else:
            requested = normal_style
        use_style = requested if requested in style_names else normal_style
        if use_style not in style_names:
            use_style = None
        paragraph = document.add_paragraph(text, style=use_style)
        cursor.addnext(paragraph._p)
        cursor = paragraph._p
        count += 1
    return count


def fill(template: Path, source: Path, mapping_path: Path, output: Path) -> list[str]:
    document = Document(template)
    originals = list(document.paragraphs)
    try:
        mapping_raw = json.loads(mapping_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError("template map is not valid JSON") from exc
    mapping = validate_map(mapping_raw, len(originals), len(document.tables))
    try:
        paper = parse_markdown(source.read_text(encoding="utf-8"))
    except (OSError, UnicodeError) as exc:
        raise ValueError("source Markdown is unreadable") from exc

    logs: list[str] = []
    delete_indices = set(mapping["delete_paragraph_indices"])
    title_index = mapping["title_anchor_para_idx"]
    abstract_index = mapping["abstract_anchor_para_idx"]
    body_index = mapping["body_anchor_para_idx"]
    references_index = mapping["references_anchor_para_idx"]
    appendix_index = mapping.get("appendix_anchor_para_idx")
    abstract_style = _sample_style(
        originals, abstract_index, body_index, delete_indices, _style_name(originals[abstract_index])
    )
    body_style = _sample_style(
        originals, body_index, references_index, delete_indices, _style_name(originals[body_index])
    )
    references_end = appendix_index if appendix_index is not None else len(originals)
    references_style = _sample_style(
        originals,
        references_index,
        references_end,
        delete_indices,
        body_style,
    )
    appendix_style = (
        _sample_style(
            originals,
            appendix_index,
            len(originals),
            delete_indices,
            body_style,
        )
        if appendix_index is not None
        else body_style
    )
    title_anchor = originals[title_index]
    title_anchor.text = paper.title
    logs.append(f"Title replaced via template_map: '{paper.title}'")

    abstract_elements = [(text, "Normal") for text in paper.abstract]
    if paper.keywords:
        abstract_elements.append((paper.keywords, "Normal"))
    _insert_elements(
        document,
        originals[abstract_index],
        abstract_elements,
        normal_style=abstract_style,
        heading_style=_style_name(originals[abstract_index]),
    )
    logs.append(f"Abstract ({len(paper.abstract)} paras) + keywords inserted")

    body_anchor = originals[body_index]
    body_count = _insert_elements(
        document,
        body_anchor,
        paper.body,
        normal_style=body_style,
        heading_style=_style_name(body_anchor),
    )
    mode = mapping["body_anchor_mode"]
    if mode == "replace":
        body_anchor.text = ""
    elif mode == "delete":
        _remove_paragraph(body_anchor)
    logs.append(f"Body ({body_count} elements) inserted (anchor_mode={mode})")

    reference_elements = [(text, "Normal") for text in paper.references]
    _insert_elements(
        document,
        originals[references_index],
        reference_elements,
        normal_style=references_style,
        heading_style=_style_name(originals[references_index]),
    )
    logs.append(f"References ({len(reference_elements)} items) inserted")

    if appendix_index is not None:
        appendix_count = _insert_elements(
            document,
            originals[appendix_index],
            paper.appendix,
            normal_style=appendix_style,
            heading_style=_style_name(originals[appendix_index]),
        )
        logs.append(f"Appendix ({appendix_count} elements) inserted")

    for index in mapping["delete_paragraph_indices"]:
        _remove_paragraph(originals[index])

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{output.stem}-", suffix=".docx", dir=output.parent, delete=False
    ) as handle:
        temporary_output = Path(handle.name)
    try:
        document.save(temporary_output)
        os.replace(temporary_output, output)
    finally:
        temporary_output.unlink(missing_ok=True)
    return logs


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True)
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--workspace", default=".")
    parser.add_argument("--map", dest="mapping", default="_template_map.json")
    args = parser.parse_args(sys.argv[1:] if argv is None else argv)
    try:
        workspace = Path(args.workspace).expanduser().resolve()
        if not workspace.is_dir() or workspace.is_symlink():
            raise ValueError("workspace must be a regular directory")
        template = _safe_path(args.template, workspace, must_exist=True, label="template")
        source = _safe_path(args.source, workspace, must_exist=True, label="source")
        mapping = _safe_path(args.mapping, workspace, must_exist=True, label="template map")
        output = _safe_path(args.output, workspace, must_exist=False, label="output")
        if template.suffix.lower() not in {".docx", ".dotx"}:
            raise ValueError("template must be DOCX or DOTX")
        if source.suffix.lower() not in {".md", ".markdown", ".txt"}:
            raise ValueError("source must be Markdown or text")
        if output.suffix.lower() != ".docx":
            raise ValueError("output must end with .docx")
        if output in {template, source, mapping}:
            raise ValueError("output must differ from every input")
        for message in fill(template, source, mapping, output):
            print(message)
        return 0
    except Exception as exc:
        print(f"DOCX_TEMPLATE_FILL_ERROR: {str(exc)[:500]}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
