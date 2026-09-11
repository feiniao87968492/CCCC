#!/usr/bin/env python3
"""Export a deterministic, zero-based structure map for a DOCX template."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document


def _safe_path(value: str, workspace: Path, *, must_exist: bool, label: str) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = workspace / path
    unresolved = path.absolute()
    resolved = path.resolve(strict=False)
    try:
        resolved.relative_to(workspace)
    except ValueError as exc:
        raise ValueError(f"{label} escapes the workspace") from exc
    current = unresolved
    while current != workspace and current.parent != current:
        if current.exists() and current.is_symlink():
            raise ValueError(f"{label} cannot use symlinks")
        current = current.parent
    if must_exist and not resolved.is_file():
        raise ValueError(f"{label} must be an existing file")
    return resolved


def analyze(template: Path) -> dict[str, object]:
    document = Document(template)
    paragraphs = [
        {
            "index": index,
            "text": paragraph.text,
            "style": paragraph.style.name if paragraph.style is not None else "",
        }
        for index, paragraph in enumerate(document.paragraphs)
    ]
    tables = []
    for index, table in enumerate(document.tables):
        cells = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append(
            {
                "index": index,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "cells": cells,
            }
        )
    language = "zh" if any("\u4e00" <= char <= "\u9fff" for item in paragraphs for char in item["text"]) else "en"
    return {
        "version": 1,
        "index_base": 0,
        "language": language,
        "paragraphs": paragraphs,
        "tables": tables,
    }


def _summary(payload: dict[str, object]) -> str:
    lines = [f"=== 模板结构清单（语言：{payload['language']}，段落索引从 0 开始）==="]
    for paragraph in payload["paragraphs"]:
        lines.append(
            f"[段 {paragraph['index']:3d}] {paragraph['style'] or '-':<18} {paragraph['text']}"
        )
    for table in payload["tables"]:
        preview = " | ".join(table["cells"][0]) if table["cells"] else ""
        lines.append(
            f"[表 {table['index']:3d}] {table['rows']}x{table['columns']} {preview}"
        )
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--summary", required=True)
    args = parser.parse_args(sys.argv[1:] if argv is None else argv)
    workspace = Path.cwd().resolve()
    try:
        template = _safe_path(args.template, workspace, must_exist=True, label="template")
        if template.suffix.lower() not in {".docx", ".dotx"}:
            raise ValueError("template must be DOCX or DOTX")
        output = _safe_path(args.output, workspace, must_exist=False, label="output")
        summary = _safe_path(args.summary, workspace, must_exist=False, label="summary")
        if len({template, output, summary}) != 3:
            raise ValueError("template, output, and summary must be different files")
        output.parent.mkdir(parents=True, exist_ok=True)
        summary.parent.mkdir(parents=True, exist_ok=True)
        payload = analyze(template)
        output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        summary.write_text(_summary(payload), encoding="utf-8")
        print(f"Analyzed {len(payload['paragraphs'])} paragraphs and {len(payload['tables'])} tables")
        return 0
    except Exception as exc:
        print(f"DOCX_TEMPLATE_ANALYZE_ERROR: {str(exc)[:500]}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
